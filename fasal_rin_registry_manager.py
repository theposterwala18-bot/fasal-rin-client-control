from __future__ import annotations

import csv
import os
import re
import shutil
import subprocess
import sys
import tkinter as tk
from datetime import date, datetime, timedelta
from pathlib import Path
from tkinter import filedialog, messagebox, ttk


ROOT = Path(__file__).resolve().parent
CSV_PATH = ROOT / "fasal_rin_registry.csv"
BUILD_SCRIPT = ROOT / "build_fasal_rin_github_files.py"
PUBLISH_ROOT = ROOT / "publish"
PUBLISH_DIR = PUBLISH_ROOT / "licenses"
BACKUP_DIR = ROOT / "registry_backups"
DAILY_BACKUP_DIR = BACKUP_DIR / "daily"
EXPORT_DIR = ROOT / "exports"

HEADERS = [
    "sr_no",
    "mobile_user_id",
    "device_id",
    "client_name",
    "pacs_bank_name",
    "district",
    "branch_name",
    "status",
    "valid_from",
    "expires_at",
    "amount_paid",
    "balance_due",
    "blocked_reason",
    "notes",
]
STATUS_VALUES = ("paid", "demo", "trial", "partial", "blocked")
LABELS = {
    "sr_no": "Sr No",
    "mobile_user_id": "Mobile No./User ID *",
    "device_id": "Device ID",
    "client_name": "Client Name *",
    "pacs_bank_name": "PACS / Bank Name *",
    "district": "District",
    "branch_name": "Branch Name",
    "status": "Status *",
    "valid_from": "Valid From * (DD-MM-YYYY)",
    "expires_at": "Expires At * (DD-MM-YYYY)",
    "amount_paid": "Amount Paid",
    "balance_due": "Balance Due",
    "blocked_reason": "Blocked Reason",
    "notes": "Notes",
}


def clean(value: object) -> str:
    return str(value or "").strip()


def identity(value: object) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "", clean(value)).upper()


def parse_date_value(value: str) -> date:
    text = clean(value)
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Date format sahi nahi: {value}. DD-MM-YYYY use karo.")


def parse_date(value: str) -> str:
    return parse_date_value(value).strftime("%d-%m-%Y")


def number_text(value: str) -> str:
    text = re.sub(r"[^0-9.\-]+", "", clean(value)) or "0"
    number = float(text)
    return str(int(number)) if number.is_integer() else f"{number:.2f}"


class FasalRinRegistryManager:
    def __init__(self) -> None:
        self.root = tk.Tk()
        self.root.title("Fasal Rin Subscription Registry")
        self.root.geometry("1380x810")
        self.root.minsize(1160, 700)
        self.root.configure(bg="#edf2f0")
        self.rows: list[dict[str, str]] = []
        self.selected_id = ""
        self.vars = {name: tk.StringVar() for name in HEADERS}
        self.search_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Ready")
        self._build_ui()
        self._ensure_daily_backup()
        self._load_rows()
        self._clear_form()
        self._schedule_daily_backup()
        self.root.mainloop()

    def _build_ui(self) -> None:
        header = tk.Frame(self.root, bg="#174f3b", height=72)
        header.pack(fill="x")
        header.pack_propagate(False)
        tk.Label(
            header,
            text="Fasal Rin Subscription Registry",
            bg="#174f3b",
            fg="white",
            font=("Segoe UI Semibold", 20),
        ).pack(side="left", padx=22, pady=17)
        tk.Label(
            header,
            text="Separate product subscription | Standalone + Integrated",
            bg="#174f3b",
            fg="#cce4db",
            font=("Segoe UI", 10),
        ).pack(side="right", padx=22)

        body = tk.Frame(self.root, bg="#edf2f0")
        body.pack(fill="both", expand=True, padx=14, pady=14)
        left = tk.Frame(body, bg="white", highlightbackground="#c7d4cf", highlightthickness=1)
        right = tk.Frame(body, bg="white", highlightbackground="#c7d4cf", highlightthickness=1)
        left.pack(side="left", fill="both", expand=True, padx=(0, 10))
        right.pack(side="left", fill="both", expand=True)
        self._build_clients(left)
        self._build_form(right)
        tk.Label(
            self.root,
            textvariable=self.status_var,
            anchor="w",
            bg="#edf2f0",
            fg="#31443d",
            font=("Segoe UI", 9),
        ).pack(fill="x", padx=14, pady=(0, 8))

    def _build_clients(self, parent: tk.Frame) -> None:
        tk.Label(parent, text="Clients", bg="white", fg="#102a43", font=("Segoe UI Semibold", 13)).pack(anchor="w", padx=14, pady=(14, 8))
        search = tk.Frame(parent, bg="white")
        search.pack(fill="x", padx=14, pady=(0, 8))
        entry = ttk.Entry(search, textvariable=self.search_var)
        entry.pack(side="left", fill="x", expand=True)
        entry.bind("<KeyRelease>", lambda _event: self._refresh_tree())
        ttk.Button(search, text="Search", command=self._refresh_tree).pack(side="left", padx=(6, 0))

        columns = ("sr", "user", "client", "pacs", "status", "expiry")
        self.tree = ttk.Treeview(parent, columns=columns, show="headings", selectmode="browse")
        labels = ("Sr", "Mobile / User ID", "Client", "PACS / Bank", "Status", "Expiry")
        widths = (45, 130, 170, 190, 75, 95)
        for name, label, width in zip(columns, labels, widths):
            self.tree.heading(name, text=label)
            self.tree.column(name, width=width, minwidth=40, anchor="w")
        scroll = ttk.Scrollbar(parent, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y", pady=(0, 14), padx=(0, 10))
        self.tree.pack(fill="both", expand=True, padx=(14, 0), pady=(0, 14))
        self.tree.bind("<<TreeviewSelect>>", self._select_row)

    def _build_form(self, parent: tk.Frame) -> None:
        tk.Label(parent, text="Fasal Rin Client Form", bg="white", fg="#102a43", font=("Segoe UI Semibold", 13)).pack(anchor="w", padx=14, pady=(14, 8))
        form = tk.Frame(parent, bg="white")
        form.pack(fill="x", padx=14)

        pairs = [
            ("sr_no", "mobile_user_id"),
            ("device_id", "client_name"),
            ("pacs_bank_name", "district"),
            ("branch_name", "status"),
            ("valid_from", "expires_at"),
            ("amount_paid", "balance_due"),
            ("blocked_reason", "notes"),
        ]
        for row_index, pair in enumerate(pairs):
            for pair_index, field in enumerate(pair):
                label_col = pair_index * 2
                input_col = label_col + 1
                tk.Label(
                    form,
                    text=LABELS[field],
                    bg="white",
                    fg="#243b53",
                    font=("Segoe UI Semibold", 9),
                    anchor="w",
                ).grid(row=row_index, column=label_col, sticky="w", padx=(0, 8), pady=7)
                if field == "status":
                    widget = ttk.Combobox(form, textvariable=self.vars[field], values=STATUS_VALUES, state="readonly")
                else:
                    widget = ttk.Entry(form, textvariable=self.vars[field])
                    if field == "sr_no":
                        widget.configure(state="readonly")
                widget.grid(row=row_index, column=input_col, sticky="ew", padx=(0, 18), pady=7, ipady=4)
        form.columnconfigure(1, weight=1)
        form.columnconfigure(3, weight=1)

        tk.Label(
            parent,
            text="Mobile No./User ID unique key hai. Public GitHub files vich raw client details nahi jandiyan; signed hashed license file hi banti hai.",
            bg="white",
            fg="#5b6b7c",
            font=("Segoe UI", 9),
            wraplength=620,
            justify="left",
        ).pack(anchor="w", padx=14, pady=(14, 8))

        top = tk.Frame(parent, bg="white")
        top.pack(fill="x", padx=14, pady=(4, 4))
        bottom = tk.Frame(parent, bg="white")
        bottom.pack(fill="x", padx=14, pady=(4, 12))
        self._button(top, "New / Clear", self._clear_form, "#52616b")
        self._button(top, "Save New", self._save_new, "#18794e")
        self._button(top, "Update Selected", self._update_selected, "#1769aa")
        self._button(bottom, "Export Excel", self._export_excel, "#087f8c")
        self._button(bottom, "Import Excel", self._import_excel, "#7c3aed")
        self._button(bottom, "Build GitHub Files", self._build_files, "#d97706")
        self._button(bottom, "Upload GitHub Now", self._upload_github, "#b42318")
        self._button(bottom, "Open Upload Folder", self._open_publish, "#51446e")

    @staticmethod
    def _button(parent: tk.Frame, text: str, command, color: str) -> None:
        tk.Button(
            parent,
            text=text,
            command=command,
            bg=color,
            fg="white",
            activebackground=color,
            activeforeground="white",
            relief="flat",
            cursor="hand2",
            font=("Segoe UI Semibold", 9),
            padx=12,
            pady=8,
        ).pack(side="left", padx=(0, 7))

    def _load_rows(self) -> None:
        self.rows = []
        dates_updated = False
        if CSV_PATH.is_file():
            with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
                for raw in csv.DictReader(handle):
                    row = {name: clean(raw.get(name)) for name in HEADERS}
                    for field_name in ("valid_from", "expires_at"):
                        if row.get(field_name):
                            formatted = parse_date(row[field_name])
                            dates_updated = dates_updated or formatted != row[field_name]
                            row[field_name] = formatted
                    self.rows.append(row)
        if dates_updated:
            self._write_rows()
        self._refresh_tree()

    def _refresh_tree(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)
        query = clean(self.search_var.get()).lower()
        for row in self.rows:
            haystack = " ".join(row.values()).lower()
            if query and query not in haystack:
                continue
            key = identity(row.get("mobile_user_id"))
            self.tree.insert(
                "",
                "end",
                iid=key,
                values=(
                    row.get("sr_no"),
                    row.get("mobile_user_id"),
                    row.get("client_name"),
                    row.get("pacs_bank_name"),
                    row.get("status"),
                    row.get("expires_at"),
                ),
            )

    def _select_row(self, _event=None) -> None:
        selected = self.tree.selection()
        if not selected:
            return
        self.selected_id = selected[0]
        row = next((item for item in self.rows if identity(item.get("mobile_user_id")) == self.selected_id), None)
        if not row:
            return
        for name in HEADERS:
            self.vars[name].set(row.get(name, ""))

    def _next_sr(self) -> str:
        values = []
        for row in self.rows:
            try:
                values.append(int(row.get("sr_no") or 0))
            except ValueError:
                pass
        return str(max(values, default=0) + 1)

    def _clear_form(self) -> None:
        self.selected_id = ""
        for var in self.vars.values():
            var.set("")
        today = date.today()
        self.vars["sr_no"].set(self._next_sr())
        self.vars["status"].set("paid")
        self.vars["valid_from"].set(today.strftime("%d-%m-%Y"))
        self.vars["expires_at"].set((today + timedelta(days=365)).strftime("%d-%m-%Y"))
        self.vars["amount_paid"].set("0")
        self.vars["balance_due"].set("0")
        for item in self.tree.selection():
            self.tree.selection_remove(item)

    def _collect(self) -> dict[str, str]:
        row = {name: clean(self.vars[name].get()) for name in HEADERS}
        required = ("mobile_user_id", "client_name", "pacs_bank_name", "status", "valid_from", "expires_at")
        missing = [LABELS[name] for name in required if not row.get(name)]
        if missing:
            raise ValueError("Required fields missing ne: " + ", ".join(missing))
        row["mobile_user_id"] = identity(row["mobile_user_id"])
        if not row["mobile_user_id"]:
            raise ValueError("Mobile No./User ID valid nahi hai.")
        if row["status"] not in STATUS_VALUES:
            raise ValueError("Status valid choose karo.")
        valid_from = parse_date_value(row["valid_from"])
        expires_at = parse_date_value(row["expires_at"])
        if expires_at < valid_from:
            raise ValueError("Expires At, Valid From ton pehla nahi ho sakda.")
        row["valid_from"] = valid_from.strftime("%d-%m-%Y")
        row["expires_at"] = expires_at.strftime("%d-%m-%Y")
        row["amount_paid"] = number_text(row["amount_paid"])
        row["balance_due"] = number_text(row["balance_due"])
        row["sr_no"] = row.get("sr_no") or self._next_sr()
        return row

    def _ensure_daily_backup(self, overwrite: bool = False) -> None:
        if not CSV_PATH.is_file():
            return
        DAILY_BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        target = DAILY_BACKUP_DIR / f"fasal_rin_registry_{date.today():%Y-%m-%d}.csv"
        if overwrite or not target.is_file():
            shutil.copy2(CSV_PATH, target)

    def _schedule_daily_backup(self) -> None:
        self._ensure_daily_backup()
        self.root.after(60 * 60 * 1000, self._schedule_daily_backup)

    def _write_rows(self) -> None:
        CSV_PATH.parent.mkdir(parents=True, exist_ok=True)
        if CSV_PATH.is_file():
            self._ensure_daily_backup()
            BACKUP_DIR.mkdir(parents=True, exist_ok=True)
            backup = BACKUP_DIR / f"fasal_rin_registry_{datetime.now():%Y%m%d_%H%M%S}.csv"
            shutil.copy2(CSV_PATH, backup)
        temp = CSV_PATH.with_suffix(".tmp")
        with temp.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=HEADERS)
            writer.writeheader()
            writer.writerows(self.rows)
        os.replace(temp, CSV_PATH)
        self._ensure_daily_backup(overwrite=True)

    def _save_new(self) -> None:
        try:
            row = self._collect()
            key = identity(row["mobile_user_id"])
            if any(identity(item.get("mobile_user_id")) == key for item in self.rows):
                raise ValueError("Eh Mobile No./User ID pehla hi maujood hai. Update Selected use karo.")
            row["sr_no"] = self._next_sr()
            self.rows.append(row)
            self._write_rows()
            self._load_rows()
            self._clear_form()
            self.status_var.set(f"Client save ho gaya: {key}")
            messagebox.showinfo("Saved", "Fasal Rin client save ho gaya.", parent=self.root)
        except Exception as exc:
            messagebox.showerror("Save Failed", str(exc), parent=self.root)

    def _update_selected(self) -> None:
        if not self.selected_id:
            messagebox.showwarning("Select Client", "Pehla client select karo.", parent=self.root)
            return
        try:
            row = self._collect()
            new_key = identity(row["mobile_user_id"])
            if new_key != self.selected_id and any(identity(item.get("mobile_user_id")) == new_key for item in self.rows):
                raise ValueError("Eh Mobile No./User ID kise hor client kol already hai.")
            index = next(i for i, item in enumerate(self.rows) if identity(item.get("mobile_user_id")) == self.selected_id)
            row["sr_no"] = self.rows[index].get("sr_no") or row["sr_no"]
            self.rows[index] = row
            self._write_rows()
            self._load_rows()
            self._clear_form()
            self.status_var.set(f"Client update ho gaya: {new_key}")
            messagebox.showinfo("Updated", "Fasal Rin client update ho gaya.", parent=self.root)
        except Exception as exc:
            messagebox.showerror("Update Failed", str(exc), parent=self.root)

    def _export_excel(self) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill

            EXPORT_DIR.mkdir(parents=True, exist_ok=True)
            default = EXPORT_DIR / f"fasal_rin_registry_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            target = filedialog.asksaveasfilename(parent=self.root, initialfile=default.name, initialdir=str(EXPORT_DIR), defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
            if not target:
                return
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Fasal Rin Clients"
            sheet.append([LABELS.get(name, name) for name in HEADERS])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="174F3B")
            for row in self.rows:
                sheet.append([row.get(name, "") for name in HEADERS])
            sheet.freeze_panes = "A2"
            workbook.save(target)
            self.status_var.set(f"Excel export ready: {target}")
        except Exception as exc:
            messagebox.showerror("Export Failed", str(exc), parent=self.root)

    def _import_excel(self) -> None:
        try:
            from openpyxl import load_workbook

            source = filedialog.askopenfilename(parent=self.root, filetypes=[("Excel", "*.xlsx")])
            if not source:
                return
            workbook = load_workbook(source, data_only=True)
            sheet = workbook.active
            raw_headers = [clean(cell.value) for cell in sheet[1]]
            reverse_labels = {label.lower(): name for name, label in LABELS.items()}
            mapped = [reverse_labels.get(value.lower(), value.lower().replace(" ", "_")) for value in raw_headers]
            imported = []
            seen = {identity(row.get("mobile_user_id")) for row in self.rows}
            for values in sheet.iter_rows(min_row=2, values_only=True):
                row = {name: "" for name in HEADERS}
                for index, value in enumerate(values):
                    if index < len(mapped) and mapped[index] in row:
                        row[mapped[index]] = clean(value)
                key = identity(row.get("mobile_user_id"))
                if not key or key in seen:
                    continue
                row["mobile_user_id"] = key
                row["sr_no"] = str(len(self.rows) + len(imported) + 1)
                row["status"] = row.get("status") or "paid"
                row["valid_from"] = parse_date(row.get("valid_from"))
                row["expires_at"] = parse_date(row.get("expires_at"))
                row["amount_paid"] = number_text(row.get("amount_paid"))
                row["balance_due"] = number_text(row.get("balance_due"))
                imported.append(row)
                seen.add(key)
            workbook.close()
            if not imported:
                raise ValueError("Koi nava valid client import nahi hoya.")
            self.rows.extend(imported)
            self._write_rows()
            self._load_rows()
            self.status_var.set(f"Excel import complete: {len(imported)} new client(s).")
            messagebox.showinfo("Import Complete", f"{len(imported)} client(s) import ho gaye.", parent=self.root)
        except Exception as exc:
            messagebox.showerror("Import Failed", str(exc), parent=self.root)

    def _build_files(self, show_message: bool = True) -> None:
        subprocess.run([sys.executable, str(BUILD_SCRIPT)], cwd=str(ROOT), check=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        self.status_var.set("Signed Fasal Rin GitHub files ready ne.")
        if show_message:
            messagebox.showinfo("Build Complete", "Signed GitHub files ready ne.", parent=self.root)

    def _upload_github(self) -> None:
        try:
            self._build_files(show_message=False)
            manifest = "licenses/Fasal_Rin_subscription_manifest.json"
            subprocess.run(["git", "add", manifest], cwd=str(ROOT), check=True)
            status = subprocess.run(["git", "status", "--porcelain", "--", manifest], cwd=str(ROOT), check=True, capture_output=True, text=True).stdout.strip()
            if not status:
                messagebox.showinfo("GitHub", "GitHub files already latest ne.", parent=self.root)
                return
            commit_message = f"Update Fasal Rin subscriptions {datetime.now():%Y-%m-%d %H:%M}"
            subprocess.run(["git", "commit", "-m", commit_message], cwd=str(ROOT), check=True)
            subprocess.run(["git", "push", "origin", "main"], cwd=str(ROOT), check=True)
            self.status_var.set("Fasal Rin subscription files GitHub te upload ho gayiyan.")
            messagebox.showinfo("Upload Complete", "GitHub upload complete ho gayi.", parent=self.root)
        except Exception as exc:
            messagebox.showerror("Upload Failed", str(exc), parent=self.root)

    def _open_publish(self) -> None:
        try:
            self._build_files(show_message=False)
            os.startfile(str(PUBLISH_ROOT))
            self.status_var.set("Manual GitHub upload folder khul gayi.")
        except Exception as exc:
            messagebox.showerror("Open Failed", str(exc), parent=self.root)


if __name__ == "__main__":
    FasalRinRegistryManager()
